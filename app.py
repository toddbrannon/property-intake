"""
Web front end for the property-intake pipeline.  Runs on Render (or anywhere).

    python app.py                       # local test at http://localhost:8000
    gunicorn app:app --bind 0.0.0.0:$PORT   # what Render runs

What it does:
  - serves one page: upload a PDF, see processed documents, work the review queue, download master.xlsx
  - starts the folder watcher (watch.py) in a background thread so uploads get processed automatically
  - optional password: set APP_PASSWORD in the environment and the page asks for it

All files live under the folders in .env / environment variables (on Render: the persistent disk).
"""
import os, csv, json, threading, time, io
from pathlib import Path
from functools import wraps
from flask import Flask, request, redirect, url_for, send_file, render_template_string, Response, abort

import extract, watch

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024  # 25 MB per upload

INBOX, DONE, FAILED, OUT = watch.INBOX, watch.DONE, watch.FAILED, extract.OUT_DIR
for d in (INBOX, DONE, FAILED, OUT):
    d.mkdir(parents=True, exist_ok=True)
PASSWORD = os.environ.get("APP_PASSWORD", "")

# ---------------------------------------------------------------- background watcher
_fields = extract.load_fields()
_started = False
def start_watcher():
    global _started
    if _started:
        return
    _started = True
    def loop():
        extract.log("web watcher thread started")
        while True:
            try:
                watch.run_once(_fields)
            except Exception as e:
                extract.log(f"watcher error: {e}")
            time.sleep(int(os.environ.get("POLL_SECONDS", "20")))
    threading.Thread(target=loop, daemon=True).start()
start_watcher()

# ---------------------------------------------------------------- optional password (login form + cookie)
import hashlib, secrets
from flask import session
app.secret_key = os.environ.get("SECRET_KEY") or hashlib.sha256(("intake:" + PASSWORD).encode()).hexdigest()

LOGIN = """<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Property Intake - sign in</title>
<style>body{margin:0;background:#f3f4f7;font:15px system-ui,sans-serif;color:#1c2030;display:flex;min-height:100vh;align-items:center;justify-content:center}
form{background:#fff;border:1px solid #d3d7e1;padding:28px 32px;min-width:280px}h1{font-size:18px;margin:0 0 14px}
input{font:inherit;padding:8px 10px;border:1px solid #d3d7e1;border-radius:3px;width:100%;box-sizing:border-box;margin-bottom:12px}
button{font:inherit;padding:8px 16px;border:0;background:#2b3f7a;color:#fff;border-radius:3px;cursor:pointer}.err{color:#b8321c;margin-bottom:10px}</style></head>
<body><form method="post"><h1>Property Intake</h1>{% if err %}<div class="err">{{ err }}</div>{% endif %}
<input type="password" name="password" placeholder="Password" autofocus><button>Sign in</button></form></body></html>"""

@app.route("/login", methods=["GET", "POST"])
def login():
    if not PASSWORD:
        return redirect(url_for("index"))
    err = None
    if request.method == "POST":
        if secrets.compare_digest(request.form.get("password", ""), PASSWORD):
            session["ok"] = True
            return redirect(url_for("index"))
        err = "Wrong password."
    return render_template_string(LOGIN, err=err)

@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))

def protected(f):
    @wraps(f)
    def inner(*a, **kw):
        if PASSWORD and not session.get("ok"):
            return redirect(url_for("login"))
        return f(*a, **kw)
    return inner

# ---------------------------------------------------------------- data helpers
def load_records():
    recs = []
    for p in sorted(OUT.glob("*.json"), key=lambda p: p.stat().st_mtime, reverse=True):
        try:
            recs.append(json.loads(p.read_text(encoding="utf-8")))
        except Exception:
            pass
    return recs

def load_review():
    rp = OUT / "review.csv"
    if not rp.exists():
        return []
    with rp.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))

def save_review(rows):
    rp = OUT / "review.csv"
    cols = extract.REVIEW_COLS
    with rp.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore"); w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, "") for k in cols})

def apply_fix_to_master(file, field, value):
    """Write a corrected value into master.xlsx and, if connected, the Google Sheet."""
    try:
        v = float(value) if value.replace(".", "", 1).replace("-", "", 1).isdigit() else value
    except Exception:
        v = value
    if field in _fields:   # mirror the correction to the Google Sheet, if one is connected
        extract.sheets_post({"action": "update", "file": file,
                             "column": f'{_fields[field]["ja"]} / {_fields[field]["en"]}', "value": v})
    from openpyxl import load_workbook
    mp = OUT / "master.xlsx"
    if not mp.exists():
        return
    wb = load_workbook(mp); ws = wb.active
    header = [c.value for c in ws[1]]
    label = _fields[field]["ja"] if field in _fields else None
    if label not in header:
        return
    col = header.index(label) + 1
    for row in ws.iter_rows(min_row=2):
        if row[0].value == file:
            row[col - 1].value = v
    wb.save(mp)

# ---------------------------------------------------------------- "what Claude added" for each document
CONVERSION_HINTS = ("万円", "億", "坪", "平成", "令和", "昭和", "converted")

def value_added(rec):
    """Facts computed from the stored record: what Claude did beyond copying text."""
    fields = rec.get("fields", {})
    n = len(fields)
    filled = sum(1 for f in fields.values() if f.get("value") is not None)
    KINDS = [(("万円", "億"), "万円/億 → yen"), (("坪",), "坪 → ㎡"), (("平成", "令和", "昭和"), "Japanese era → calendar year")]
    conv, kinds = [], []
    for k, f in fields.items():
        if f.get("value") is None or _fields.get(k, {}).get("type") not in ("integer", "number", "date"):
            continue
        src = f.get("source") or ""
        hit = [label for keys, label in KINDS if any(x in src for x in keys)]
        if hit and str(f.get("value")) not in src:   # the stored value differs from the printed text: a real conversion
            conv.append(k); kinds += [h for h in hit if h not in kinds]
    risks = (fields.get("risk_flags") or {}).get("value") or []
    blanks_req = [k for k, f in fields.items() if f.get("value") is None and _fields.get(k, {}).get("required")]
    blanks = [k for k, f in fields.items() if f.get("value") is None]
    yld = next((it for it in rec.get("review", []) if it.get("field") == "gross_yield_pct" and "income" in it.get("issue", "")), None)
    has_yield_inputs = all((fields.get(k) or {}).get("value") for k in ("price_jpy", "annual_income_jpy", "gross_yield_pct"))
    out = [f"Read the Japanese PDF and filled {filled} of {n} fields, each with the page and the original Japanese text it came from."]
    if conv:
        names = ", ".join(_fields[k]["en"] for k in conv[:4] if k in _fields) + ("…" if len(conv) > 4 else "")
        out.append(f"Converted {len(conv)} value{'s' if len(conv) != 1 else ''} into standard units ({names}): {', '.join(kinds)}.")
    if risks:
        out.append(f"Pulled {len(risks)} buyer risk{'s' if len(risks) != 1 else ''} out of the remarks and fine print into their own column.")
    if blanks:
        req = f", including {len(blanks_req)} required field{'s' if len(blanks_req) != 1 else ''} sent to a person" if blanks_req else ""
        out.append(f"Left {len(blanks)} field{'s' if len(blanks) != 1 else ''} blank because the document doesn't state them{req}. Nothing was guessed.")
    if yld:
        out.append("The yield check failed: the broker's printed yield doesn't match income ÷ price. A person re-typing by hand would usually copy the printed figure without noticing.")
    elif has_yield_inputs:
        out.append("The yield check passed: the printed yield agrees with income ÷ price.")
    if rec.get("summary_en"):
        out.append("Wrote the English summary above for colleagues who don't read Japanese.")
    return out

# ---------------------------------------------------------------- page
PAGE = """<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Property Intake</title>
<style>
 :root{--bg:#f3f4f7;--fg:#1c2030;--muted:#5b6275;--line:#d3d7e1;--accent:#2b3f7a;--ok:#1f7a4d;--oks:#d9efe3;--flag:#b8321c;--flags:#f8ddd6;--warn:#9a6a00;--warns:#f6ecc8}
 body{margin:0;background:var(--bg);color:var(--fg);font:14px/1.5 system-ui,-apple-system,"Segoe UI",Roboto,"Hiragino Sans","Noto Sans JP",sans-serif}
 .wrap{max-width:1100px;margin:0 auto;padding:24px 16px 60px}
 h1{font-size:22px;margin:0 0 4px} h2{font-size:16px;margin:28px 0 10px} .muted{color:var(--muted)}
 .card{background:#fff;border:1px solid var(--line);padding:14px 16px;margin-top:10px}
 form.up{display:flex;gap:10px;flex-wrap:wrap;align-items:center}
 .btn{font:inherit;padding:7px 14px;border:1px solid var(--accent);background:var(--accent);color:#fff;cursor:pointer;border-radius:3px}
 .btn.sec{background:#fff;color:var(--fg);border-color:var(--line)}
 table{border-collapse:collapse;width:100%;font-size:13px} th,td{text-align:left;padding:6px 8px;border-bottom:1px solid var(--line);vertical-align:top}
 th{font-size:11px;text-transform:uppercase;letter-spacing:.06em;color:var(--muted)}
 .pill{display:inline-block;padding:1px 8px;border-radius:3px;font-size:12px;font-weight:500}
 .ok{background:var(--oks);color:var(--ok)} .flag{background:var(--flags);color:var(--flag)} .warn{background:var(--warns);color:var(--warn)}
 .src{color:var(--muted);font-size:12px} input[type=text]{font:inherit;padding:4px 6px;border:1px solid var(--line);border-radius:3px;width:180px}
 .inline{display:inline-flex;gap:6px;align-items:center;flex-wrap:wrap} .msg{background:var(--oks);color:var(--ok);padding:8px 12px;margin-top:10px}
 .tw{overflow-x:auto} details summary{cursor:pointer;color:var(--accent)}
 .how{background:#fff;border:1px solid var(--line);border-left:4px solid var(--accent);padding:14px 18px;margin-top:14px}
 .how summary{font-weight:600;color:var(--accent);font-size:15px}
 .steps{display:grid;grid-template-columns:repeat(4,1fr);gap:12px;margin-top:12px} @media(max-width:760px){.steps{grid-template-columns:1fr 1fr}}
 .step{background:var(--bg);padding:10px 12px;border-radius:4px} .step b{display:block;margin-bottom:2px} .step.ai{background:#e4e9f6;outline:1px solid #b9c4e3}
 .tag{display:inline-block;font-size:11px;font-weight:600;letter-spacing:.04em;text-transform:uppercase;padding:1px 6px;border-radius:3px;margin-right:4px}
 .tag.ai{background:#e4e9f6;color:var(--accent)} .tag.rule{background:#eceae4;color:#5a4f2e}
 .why{font-size:13px;color:var(--muted);margin-top:4px;max-width:70ch} .cnote{font-size:13px;margin-top:4px;padding:6px 10px;background:#f6f7fa;border-left:3px solid #b9c4e3;max-width:70ch}
 .meta{display:flex;flex-wrap:wrap;gap:6px 14px;font-size:12.5px;color:var(--muted);margin-top:8px;padding-top:8px;border-top:1px dashed var(--line)}
 .meta b{color:var(--fg);font-weight:600} .summary{margin-top:8px;max-width:80ch}
 .fnote{display:block;font-size:12px;color:#4a5168;font-style:italic;margin-top:2px}
 .va{margin-top:10px;background:#f4f6fb;border:1px solid #d5dcef;border-radius:4px;padding:10px 14px}
 .va h4{margin:0 0 6px;font-size:13px;color:var(--accent);text-transform:uppercase;letter-spacing:.05em}
 .va ul{margin:0;padding-left:18px} .va li{margin:3px 0;max-width:90ch} .va .lbl{font-size:12px;color:var(--muted);margin:8px 0 2px;font-weight:600}
 .va .saved{margin-top:8px;font-size:12.5px;color:var(--ok);font-weight:600}
</style></head><body><div class="wrap">
<h1>Property Intake</h1>
<div class="muted">Upload a property PDF. It is read, checked, and added to the master sheet; unsure fields come to the review queue. Read by: <b>{{ mode }}</b>.</div>
{% if msg %}<div class="msg">{{ msg }}</div>{% endif %}

<div class="card"><form class="up" method="post" action="{{ url_for('upload') }}" enctype="multipart/form-data">
 <input type="file" name="pdf" accept="application/pdf" required>
 <button class="btn" type="submit">Upload PDF</button>
 <span class="muted">Waiting in inbox: {{ inbox|length }} · Failed: {{ failed|length }}</span>
 {% if sheet_url %}<a class="btn" style="background:#1e8e3e;border-color:#1e8e3e" href="{{ sheet_url }}" target="_blank" rel="noopener">Open Google Sheet</a>{% endif %}
 <a class="btn sec" href="{{ url_for('download') }}">Download master.xlsx</a>
 <a class="btn sec" href="{{ url_for('index') }}">Refresh</a>{% if pw %} <a class="btn sec" href="{{ url_for('logout') }}">Sign out</a>{% endif %}
</form>
{% if not has_key %}<div class="muted" style="margin-top:8px">No API key set: the three sample documents work; other PDFs will go to Failed until <code>ANTHROPIC_API_KEY</code> is added in Render → Environment.</div>{% endif %}
</div>

<details class="how" open><summary>How this works: what Claude does, and what it doesn't</summary>
<div class="steps">
 <div class="step"><b>1 · Upload</b>The PDF is saved to the inbox. A background worker picks it up within about 20 seconds.</div>
 <div class="step ai"><b>2 · Claude reads it</b>One call to the Claude API ({{ model }}). Claude reads the Japanese PDF directly, no translation step, and returns all {{ nfields }} fields. For each one it gives the value, the page, the exact Japanese text it read, a confidence level, and a note when it's unsure. It also writes an English summary.</div>
 <div class="step"><b>3 · The script checks Claude</b>Plain Python, not AI: types and units, required fields, allowed categories, and arithmetic (yield = income ÷ price). Claude is never trusted on its own.</div>
 <div class="step"><b>4 · A person decides the rest</b>Only fields that Claude was unsure of, or that a rule caught, reach the review queue. Everything else goes straight to the spreadsheet.</div>
</div>
<p class="why" style="margin-top:10px"><span class="tag ai">Claude</span> marks items Claude itself flagged as uncertain. <span class="tag rule">Rule</span> marks items caught by a check in the script. Confidence: <span class="pill ok">high</span> read directly · <span class="pill warn">medium</span> converted (坪→㎡, 平成→year, 万円→yen) · <span class="pill flag">low</span> ambiguous, sent to review. The ground rule in Claude's instructions: <b>if the document doesn't say it, leave it blank and never guess.</b></p>
</details>

<h2>Review queue <span class="muted">({{ open_items|length }} open)</span></h2>
{% if not open_items %}<div class="card muted">Nothing waiting. Everything processed so far was accepted automatically.</div>{% endif %}
{% for it in open_items %}
<div class="card"><form method="post" action="{{ url_for('decide') }}" class="inline">
 <input type="hidden" name="file" value="{{ it.file }}"><input type="hidden" name="field" value="{{ it.field }}">
 <div style="flex:1 1 340px"><b>{{ it.label_ja }}</b> <span class="muted">({{ it.field }})</span> · {{ it.file }}<br>
 <span class="tag {{ 'ai' if it.raised_by and 'Claude' in it.raised_by else 'rule' }}">{{ 'Claude' if it.raised_by and 'Claude' in it.raised_by else 'Rule' }}</span><span class="flag pill">{{ it.issue }}</span><br>
 <span class="src">p.{{ it.page or '–' }}  {{ it.source_text }}</span>
 {% if it.why %}<div class="why"><b>Why it's here ({{ it.raised_by }}):</b> {{ it.why }}</div>{% endif %}
 {% if it.claude_note %}<div class="cnote"><b>{{ 'The calculation' if it.raised_by == 'Arithmetic cross-check' else "Claude's note" }}:</b> {{ it.claude_note }}</div>{% endif %}</div>
 <input type="text" name="value" placeholder="corrected value">
 <button class="btn" name="decision" value="fix">Save fix</button>
 <button class="btn sec" name="decision" value="ok">OK as is</button>
 <button class="btn sec" name="decision" value="skip">Skip</button>
</form></div>
{% endfor %}

<h2>Documents <span class="muted">({{ records|length }})</span></h2>
{% for r in records %}
<div class="card"><div class="inline" style="justify-content:space-between">
 <div><b>{{ r.file }}</b> <span class="muted">{{ r.processed }}</span></div>
 <span class="pill {{ 'flag' if r.status=='needs_review' else 'ok' }}">{{ r.status.replace('_',' ') }}</span></div>
 {% if r.summary_en %}<p class="summary">{{ r.summary_en }}</p>{% endif %}
 <div class="va"><h4>What Claude added</h4>
 {% if r.findings %}<div class="lbl">Claude's findings on this document</div><ul>{% for f in r.findings %}<li>{{ f }}</li>{% endfor %}</ul>
 <div class="lbl">What it did, counted from the record</div>{% endif %}
 <ul>{% for line in va[r.file] %}<li>{{ line }}</li>{% endfor %}</ul>
 <div class="saved">Manual equivalent: about 20–35 minutes of reading, re-typing and summarizing. With Claude: {{ r.review|length }} item{{ '' if r.review|length == 1 else 's' }} for a person to confirm.</div>
 </div>
 <div class="meta">
 {% if r.read_by == 'api' and r.meta %}<span><span class="tag ai">Claude</span> <b>{{ r.meta.model }}</b> read this PDF</span><span><b>{{ '{:,}'.format(r.meta.input_tokens) }}</b> tokens in · <b>{{ '{:,}'.format(r.meta.output_tokens) }}</b> out</span><span>cost <b>${{ '%.3f'|format(r.meta.cost_usd) }}</b></span><span><b>{{ r.meta.seconds }}s</b></span>
 {% elif r.read_by == 'demo' %}<span><span class="tag rule">Recorded</span> built-in sample: Claude's reading was recorded earlier, so no API call was made</span>
 {% elif r.read_by == 'manual' %}<span><span class="tag ai">Claude</span> read by hand in the Claude app and pasted in</span>{% endif %}
 {% if r.stats %}<span>Fields: <b>{{ r.stats.high }}</b> high · <b>{{ r.stats.medium }}</b> medium · <b>{{ r.stats.low }}</b> low · <b>{{ r.stats.empty }}</b> not in document</span>{% endif %}
 <span>Sent to review: <b>{{ r.review|length }}</b></span>
 </div>
 <details><summary>{{ r.fields|length }} fields</summary><div class="tw"><table><tr><th>Field</th><th>Value</th><th>Source · page</th><th>Conf.</th></tr>
 {% for k, f in r.fields.items() %}<tr><td>{{ labels[k] }}<br><span class="src">{{ labels_en[k] }}</span></td>
 <td>{% if f.value is none %}<span class="muted">— not in document</span>{% elif f.value is iterable and f.value is not string %}{{ f.value|join('; ') }}{% else %}{{ f.value }}{% endif %}</td>
 <td class="src">p.{{ f.page or '–' }}  {{ f.source }}{% if f.note %}<span class="fnote">Claude: {{ f.note }}</span>{% endif %}</td>
 <td><span class="pill {{ 'ok' if f.confidence=='high' else ('warn' if f.confidence=='medium' else 'flag') }}">{{ f.confidence }}</span></td></tr>{% endfor %}
 </table></div></details></div>
{% endfor %}

{% if failed %}<h2>Failed <span class="muted">({{ failed|length }})</span></h2>
{% for f in failed %}<div class="card"><b>{{ f.name }}</b><br><span class="src">{{ f.err }}</span>
 <form method="post" action="{{ url_for('retry') }}" style="margin-top:6px"><input type="hidden" name="name" value="{{ f.name }}"><button class="btn sec">Retry</button></form></div>{% endfor %}{% endif %}

<p class="muted" style="margin-top:30px"><a href="{{ url_for('logtail') }}">log</a> · samples: {% for s in samples %}<a href="{{ url_for('use_sample', name=s) }}">{{ s }}</a>{% if not loop.last %}, {% endif %}{% endfor %}</p>
</div></body></html>"""

@app.route("/")
@protected
def index():
    records = load_records()
    review = load_review()
    open_items = [r for r in review if not r.get("decision")]
    open_files = {r["file"] for r in open_items}
    for rec in records:   # a document is "reviewed" once none of its items are open
        if rec.get("status") == "needs_review" and rec["file"] not in open_files:
            rec["status"] = "reviewed"
    failed = []
    for p in sorted(FAILED.glob("*.pdf")):
        e = FAILED / f"{p.stem}.error.txt"
        err = e.read_text(encoding="utf-8").strip().splitlines()[-1] if e.exists() else ""
        failed.append({"name": p.name, "err": err})
    samples = [p.name for p in sorted((Path(__file__).parent / "samples").glob("*.pdf"))]
    va = {r["file"]: value_added(r) for r in records}
    return render_template_string(PAGE, va=va, records=records, open_items=open_items, failed=failed,
        inbox=list(INBOX.glob("*.pdf")), samples=samples, msg=request.args.get("msg"),
        labels={k: v["ja"] for k, v in _fields.items()}, labels_en={k: v["en"] for k, v in _fields.items()},
        has_key=bool(os.environ.get("ANTHROPIC_API_KEY")), pw=bool(PASSWORD), model=extract.MODEL, nfields=len(_fields), sheet_url=os.environ.get("SHEETS_URL", ""),
        mode="Claude API" if os.environ.get("ANTHROPIC_API_KEY") else "samples only (no API key)")

@app.route("/upload", methods=["POST"])
@protected
def upload():
    f = request.files.get("pdf")
    if not f or not f.filename.lower().endswith(".pdf"):
        return redirect(url_for("index", msg="Please choose a PDF file."))
    name = Path(f.filename).name.replace(" ", "_")
    f.save(INBOX / name)
    return redirect(url_for("index", msg=f"Uploaded {name}. It will be processed within about 20 seconds; click Refresh."))

@app.route("/sample/<name>")
@protected
def use_sample(name):
    src = Path(__file__).parent / "samples" / Path(name).name
    if not src.exists():
        abort(404)
    (INBOX / src.name).write_bytes(src.read_bytes())
    return redirect(url_for("index", msg=f"Queued sample {src.name}."))

@app.route("/decide", methods=["POST"])
@protected
def decide():
    file, field, decision, value = request.form["file"], request.form["field"], request.form["decision"], request.form.get("value", "").strip()
    rows = load_review()
    for r in rows:
        if r["file"] == file and r["field"] == field and not r.get("decision"):
            r["decision"] = decision
            if decision == "fix":
                r["corrected_value"] = value
                apply_fix_to_master(file, field, value)
    save_review(rows)
    extract.log(f"review: {file} / {field} -> {decision} {value}")
    return redirect(url_for("index"))

@app.route("/retry", methods=["POST"])
@protected
def retry():
    p = FAILED / Path(request.form["name"]).name
    if p.exists():
        (FAILED / f"{p.stem}.error.txt").unlink(missing_ok=True)
        p.rename(INBOX / p.name)
    return redirect(url_for("index", msg="Moved back to inbox."))

@app.route("/master.xlsx")
@protected
def download():
    mp = OUT / "master.xlsx"
    if not mp.exists():
        return redirect(url_for("index", msg="No master.xlsx yet; process a document first."))
    return send_file(mp, as_attachment=True, download_name="master.xlsx")

@app.route("/log")
@protected
def logtail():
    lp = OUT / "log.txt"
    text = lp.read_text(encoding="utf-8")[-20000:] if lp.exists() else "(no log yet)"
    return Response(text, mimetype="text/plain")

@app.route("/health")
def health():
    return "ok"

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "8000")), debug=False)
